import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from tempfile import NamedTemporaryFile

async def create_supplier_requests_report_xlsx(data: list) -> str:
    """
    Создаёт xlsx-файл с отчётом по заявкам на поставщиков.
    Возвращает путь к временному файлу.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заявки на поставщиков"

    headers = [
        "ID заявки", "Дата", "Название компании", "Категория", "Подкатегория", "Статус", "Проверил"
    ]
    ws.append(headers)

    for row in data:
        ws.append([
            row.get("supplier_id"),
            row.get("date"),
            row.get("company_name"),
            row.get("main_category"),
            row.get("category"),
            row.get("status"),
            row.get("verified_by")
        ])

    # Добавляем автофильтр
    ws.auto_filter.ref = ws.dimensions

    # Стилизация таблицы
    tab = Table(displayName="RequestsTable", ref=ws.dimensions)
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws.add_table(tab)

    # Автоширина колонок
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    # Сохраняем во временный файл
    with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        return tmp.name
